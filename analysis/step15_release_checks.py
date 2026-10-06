"""v13 release checks (v10 checks plus: main-text reference count <= 50, abstract <= 250 words, TOP and IRB statements present). v10 release checks (v9 checks plus: contamination-mixture reconstruction, hypertension coloc with modelled SEs, window check, S80 Wald recomputation, parenthesis and wording checks). v9 release checks (v8 checks plus review-response checks: robust estimators, Nakao weights, Steiger trail, conditional colocalization, hypertension SE sensitivity). v8 release checks and manifest (v7 checks plus LD-panel allele-key checks, qhet objective and bootstrap checks, embedded-figure identity). Run after all analyses (before the documents are built): machine-readable PASS/FAIL checks that the
publication tables belong to one run and reconstruct from their named inputs, plus a SHA-256 manifest and a release ID derived from the hashes
of all inputs, scripts and result files. With --render (after the documents are built) it adds checks that key values in the rendered
manuscript equal the result files, and hashes the rendered outputs. The build fails if any check fails."""
import sys, os, glob, json, hashlib, re
import numpy as np, pandas as pd
from docx.oxml.ns import qn as qn_
from scipy import stats
os.makedirs('release_v13', exist_ok=True)
C = []
def add(check, ok, detail=''): C.append(dict(check=check, status='PASS' if ok else 'FAIL', detail=detail))
def ivw(bx, by, sy):
    w = 1 / sy ** 2; b = (w * bx * by).sum() / (w * bx ** 2).sum(); k = len(bx)
    rss = (w * (by - b * bx) ** 2).sum(); se = np.sqrt(max(1.0, rss / (k - 1)) / (w * bx ** 2).sum()); p = 2 * stats.t.sf(abs(b / se), k - 1)
    return b, se, p, rss, k
R = pd.read_csv('step8_mr_results_v6.csv')
key = R.direction.astype(str) + '|' + R.code + '|' + R.analysis + '|' + R.method + '|' + R.source.astype(str)
add('S11 analysis keys unique (direction|code|analysis|method|source)', not key.duplicated().any(), f'{len(R)} rows')
O = pd.read_csv('step9_mr_output_consistency_S46.csv'); add('S46: every S11 row reproduces P and CI from b and SE', (O.status == 'PASS').all() and len(O) == len(R), f'{(O.status == "PASS").sum()}/{len(O)} rows')
T = pd.read_csv('step9_reverse_reconstruction_test.csv'); add('S72: every reverse IVW result reconstructs from S13 rows (k, b, SE, P)', (T.status == 'PASS').all(), f'{len(T)} results; max rel diff {T.max_rel_diff.max():.1e}')
# forward IVW from S12 rows (primary, FinnGen) and PLINK forward inputs
S12 = pd.read_csv('step6_v4_harmonised_inputs_all.csv'); worst = 0; n = 0; okk = True
for code in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU', 'H_fg12']:
    d = S12[(S12.code == code) & (S12.file == f'forward_selected_r2_0.001_{code}.csv')]
    b, se, p, rss, k = ivw(d.bx.values, d.by.values, d.sy.values); st = R[(R.code == code) & (R.analysis == 'r2<0.001 (primary)') & (R.method == 'IVW (MRE, t)')].iloc[0]
    rel = max(abs(b - st.b) / abs(st.b), abs(se - st.se) / st.se); worst = max(worst, rel); n += 1; okk &= (k == st.nsnp) and rel < 1e-8
    d = pd.read_csv(f'mr_v6_inputs/forward_plink_{code}.csv'); b, se, p, rss, k = ivw(d.bx.values, d.by.values, d.sy.values)
    st = R[(R.code == code) & (R.analysis == 'PLINK 1.9 clumping') & (R.method == 'IVW (MRE, t)')].iloc[0]
    rel = max(abs(b - st.b) / abs(st.b), abs(se - st.se) / st.se); worst = max(worst, rel); n += 1; okk &= (k == st.nsnp) and rel < 1e-8
add('Forward IVW (primary from S12 rows; PLINK sets) reconstructs S11', okk, f'{n} results; max rel diff {worst:.1e}')
# MR-PRESSO corrected estimates at B = 1e6 reconstruct from S12 minus the listed outliers
CR = pd.read_csv('step8_presso_convergence_runs.csv'); okp = True; det = []
for code in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU']:
    r6 = CR[(CR.code == code) & (CR.B == 1000000)]; sets = r6.outliers.fillna('').map(lambda x: frozenset(x.split(';')) - {''})
    same = sets.nunique() == 1; okp &= same
    d = S12[(S12.code == code) & (S12.file == f'forward_selected_r2_0.001_{code}.csv')]; keep = ~d.rsid.isin(sets.iloc[0])
    dd = d[keep]; w = 1 / dd.sy ** 2; b = (w * dd.bx * dd.by).sum() / (w * dd.bx ** 2).sum()
    st = R[(R.code == code) & (R.method == 'MR-PRESSO (outlier-corrected)')].iloc[0]; okp &= abs(b - st.b) < 1e-10 and int(keep.sum()) == int(st.nsnp)
    det.append(f'{code}: {len(sets.iloc[0])} outliers, k {int(keep.sum())}')
add('MR-PRESSO B = 1e6: identical outlier sets over 5 seeds; corrected b reconstructs from S12', okp, '; '.join(det))
E = pd.read_csv('step9_presso_equivalence.csv'); E = E[E.check != 'summary']
add('S70: Python MR-PRESSO agrees with MRPRESSO 1.0 on shared simulated data within stated tolerances (all intermediates)', len(E) > 0 and (E.status == 'PASS').all() and E.code.nunique() == 3, f'{(E.status == "PASS").sum()}/{len(E)} checks, {E.code.nunique()} outcomes')
# MVMR: weighted regression on S20 minus S39 reproduces S21 (all three exposures)
M = pd.read_csv('mr_v4_inputs/mvmr_instruments.csv'); fb = set(pd.read_csv('step7_mvmr_fallback_variants.csv').varId); M = M[~M.varId.isin(fb)]
E21 = pd.read_csv('step9_mvmr_estimates_primary.csv'); okm = len(M) == 318; worst = 0
for oc in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU']:
    X = M[['LTL_b', 'BMI_b', 'LYM_b']].values; y = M[oc + '_b'].values; w = 1 / M[oc + '_se'].values ** 2
    beta = np.linalg.solve(X.T @ (X * w[:, None]), X.T @ (w * y))
    for j, ex in enumerate(['LTL', 'BMI', 'LYM']):
        st = E21[(E21.outcome == oc) & (E21.exposure == ex) & (E21.method == 'MV-IVW (random, t)')].iloc[0]
        rel = abs(beta[j] - st.b) / abs(st.b); worst = max(worst, rel); okm &= rel < 1e-8 and int(st.nsnp) == 318
add('S21: MV-IVW coefficients (LTL, BMI, LYM) reconstruct from S20 minus S39 (318 variants)', okm, f'max rel diff {worst:.1e}')
Q = pd.read_csv('step9_mvmr_qhet.csv'); g = Q.groupby(['outcome', 'rho'])
okq = (g.size() == 3).all() and (g.exposure.apply(lambda s: list(s) == ['LTL', 'BMI', 'LYM'])).all() and (g.qhet_estimate.nunique() == 3).all()
V = pd.read_csv('step9_qhet_vectorised_check.csv'); okq &= V.abs_diff.max() < 1e-10
add('S55: qhet_mvmr exported one row per exposure (LTL, BMI, LYM) per outcome and rho; three distinct values; vectorised objective equals reference', bool(okq), f'{len(g)} outcome x rho groups')
# variant accounting for reverse PLINK sets before/after outcome matching
RL = pd.read_csv('step8_reverse_source_ledger.csv'); okv = True; det = []
for r in RL.itertuples():
    an = 'r2<0.001 (primary)' if r.selection == 'custom' else 'PLINK 1.9 clumping'
    st = R[(R.code == f'{r.exposure}->L_codd') & (R.analysis == an) & (R.method == 'IVW (MRE, t)')].iloc[0]
    okv &= (r.instruments - r.absent - r.allele_mismatch) == st.nsnp; det.append(f'{r.exposure}/{r.selection}: {r.instruments} selected, {r.absent} without source record, k {int(st.nsnp)}')
add('Reverse variant accounting: selected - unmatched = k used', okv, '; '.join(det))
L52 = pd.read_csv('step8_plink_discrepancy_ledger.csv')
okl = L52.plink_r2_dprime_ML.notna().all() and (L52.mechanism.str.startswith('threshold') | L52.mechanism.str.startswith('selection chain')).all()
c = L52[L52.plink_clump_r2.notna()]; okl &= ((c.plink_clump_r2 - c.plink_r2_dprime_ML).abs() / c.plink_clump_r2).max() < 0.005
add('S52: every discrepant variant has a PLINK --r2 dprime value and an explained mechanism; dprime r2 matches clump r2 (<0.5%)', bool(okl), f'{len(L52)} rows')
L38 = pd.read_csv('step9_se_fallback_ledger_S38.csv')
add('S38: no reverse fallback record is marked as used; MVMR fallback variants all excluded', (L38[L38.direction == 'reverse'].status_current_release.str.startswith('superseded')).all() and set(L38[L38.direction == 'MVMR'].varId).isdisjoint(set(M.varId)), f'{len(L38)} records')

# ---- v8 checks
# LD-panel allele key: every variant written to a PLINK fileset carries its own alleles (either orientation), and no script that reads the
# LD panel uses a position-only key
okld = True; det = []
for vcf in sorted(glob.glob('plink_clump/*.vcf')):
    if os.path.basename(vcf) in ('pair.vcf',): continue
    bad = n = 0
    ids = {}
    for tag, U in [('fwd', pd.read_csv('step6_v4_forward_variant_ledger.csv'))] + [(f'ref_{e.lower()}', pd.read_csv('step3v3_reverse_variant_ledgers_all.csv').query('exposure == @e')) for e in ['SBP', 'DBP', 'Hypertension']]:
        if os.path.basename(vcf).startswith(tag): ids = dict(zip(U.rsid, U.varId))
    if not ids: continue
    for line in open(vcf):
        if line.startswith('#'): continue
        f = line.split('\t', 5); n += 1; v = ids.get(f[2])
        if v is None: bad += 1; continue
        a0, a1 = v.split(':')[2:4]
        if {f[3], f[4]} != {a0, a1}: bad += 1
    okld &= bad == 0; det.append(f'{os.path.basename(vcf)}: {n} variants, {bad} allele mismatches')
posonly = [f for f in glob.glob('*.py') if 'ld_ref' in open(f).read() and re.search(r"key = \{\(r\.chr, r\.pos\): i", open(f).read()) and not f.startswith(('step9_release', 'step10_release', 'step11_release', 'step12_release', 'step13_release', 'step14_release', 'step15_release'))]
add('LD panel: PLINK filesets carry each variant\'s own alleles; no current LD script uses a position-only key', okld and not posonly, '; '.join(det) + (f'; position-only key in {posonly}' if posonly else ''))
# every LD-panel lookup through ld_ref.match (primary clumping in both directions, MVMR, GLS set, colocalization windows) returns a record
# whose alleles are the variant's own alleles; palindromic variants are flagged by match() and handled by the documented rule
from ld_ref import load as _load, match as _match
okm2 = True; det = []
REFI = _load('ldref/eur_instruments.json')
sets_ = [('forward candidates', pd.read_csv('step6_v4_forward_variant_ledger.csv').varId), ('reverse candidates', pd.read_csv('step3v3_reverse_variant_ledgers_all.csv').varId.drop_duplicates()),
         ('MVMR instruments', pd.read_csv('mr_v4_inputs/mvmr_instruments.csv').varId)]
REFM = pd.concat([_load('ldref/eur_instruments.json'), _load('ldref/eur_mvmr_candidates.json')]).drop_duplicates(['chr', 'pos', 'ref', 'alt']).reset_index(drop=True)   # as in step6_mvmr_prep.py
for lab, vids in sets_:
    P_ = REFM if lab.startswith('MVMR') else REFI
    mm = _match(list(vids), P_); used = mm[mm.idx >= 0]
    badn = sum(1 for r in used.itertuples() if {P_.iloc[r.idx].ref, P_.iloc[r.idx].alt} != set(r.varId.split(':')[2:4]))
    okm2 &= badn == 0; det.append(f'{lab}: {len(used)} matched records, {badn} allele mismatches, {int(used.palindromic.sum())} palindromic flagged')
for loc in ['TERC', 'TERT']:
    REFR = _load(f'ldref/eur_region_{loc}.json'); vids = pd.concat([pd.read_csv(f) for f in glob.glob(f'coloc_v4_inputs/{loc}_250kb_*.csv') if not f.endswith(('_LD.csv', '_LDdos.csv'))]).varId.drop_duplicates()
    mm = _match(list(vids), REFR); used = mm[mm.idx >= 0]
    badn = sum(1 for r in used.itertuples() if {REFR.iloc[r.idx].ref, REFR.iloc[r.idx].alt} != set(r.varId.split(':')[2:4]))
    okm2 &= badn == 0; det.append(f'{loc} colocalization window: {len(used)} matched records, {badn} allele mismatches')
add('LD panel: every allele-aware lookup (clumping candidates, MVMR, colocalization windows) returns the variant\'s own alleles', okm2, '; '.join(det))
T74 = pd.read_csv('step10_ld_lookup_correction_S74.csv')
add('S74: variants affected by the former position-only lookup are exactly rs113695388 and rs147632758, and neither is in a current PLINK set', set(T74[~T74.position_only_record_correct].rsid) == {'rs113695388', 'rs147632758'} and not T74[~T74.position_only_record_correct].plink_kept_current.any(), f'{len(T74)} multi-allelic candidates')
QO = pd.read_csv('step10_qhet_objective_test.csv')
add('qhet objective: vectorised objective equals the reference PL2_MVMR (extracted from qhet_mvmr) at random coefficient vectors and tau2', QO.rel_diff.max() < 1e-10, f'{len(QO)} points; max rel diff {QO.rel_diff.max():.1e}')
okb = True; det = []
for oc in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU']:
    Rb = pd.read_csv(f'qhet_bootstrap/{oc}_rho0_replicates.csv'); q = Q[(Q.outcome == oc) & (Q.rho == 0)]
    ok_ = Rb.fit_ok == 1
    for ex in ['LTL', 'BMI', 'LYM']:
        r = q[q.exposure == ex].iloc[0]
        okb &= abs(np.quantile(Rb.loc[ok_, ex], 0.025) - r.pct_lo) < 1e-12 and abs(np.quantile(Rb.loc[ok_, ex], 0.975) - r.pct_hi) < 1e-12 and len(Rb) == r.bootstrap_B
    det.append(f'{oc}: {len(Rb)} replicates, {int((~ok_).sum())} failed, seed {int(Rb.seed.iloc[0])}')
add('S55 bootstrap intervals reconstruct from the saved replicates (qhet_bootstrap/)', bool(okb), '; '.join(det))
# no stale sources in the publication workbook builder
bs = open('build_supplement_v13.py').read()
stale = [f for f in ['step6_v4_reverse_inputs_all.csv', 'step7_se_fallback_ledger.csv', 'step7_mr_output_consistency.csv', 'step8_mvmr_v6_qhet.csv', 'step8_mvmr_v6_estimates.csv'] if f"'{f}'" in bs]
add('Workbook uses no superseded source for a current table', not stale, ', '.join(stale) or 'none')


# ---- v9 checks (pre-submission review)
RB = pd.read_csv('step11_robust_mr.csv'); okr = True; det = []
for code in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU', 'H_fg12']:
    a = RB[(RB.code == code) & (RB.method == 'IVW (MRE, t) [reference]')].iloc[0]; st = R[(R.code == code) & (R.analysis == 'r2<0.001 (primary)') & (R.method == 'IVW (MRE, t)')].iloc[0]
    okr &= abs(a.b - st.b) < 1e-10 and int(a.k) == int(st.nsnp)
    sub = RB[(RB.code == code) & (RB.weights == 'Codd 2021')]; okr &= sub.method.nunique() == 4 and sub.b.notna().all() and (sub.k == 116).all()
    det.append(f'{code}: ' + ', '.join(f'{m.split(" (")[0]} {b:.3f}' for m, b in zip(sub.method, sub.b)))
add('S80: robust estimators run on the primary 116-instrument inputs (reference IVW equals S11)', bool(okr), '; '.join(det))
okn = True
for code in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU', 'H_fg12']:
    a = pd.read_csv(f'mr_v4_inputs/forward_selected_r2_0.001_{code}.csv'); b_ = pd.read_csv(f'mr_v9_inputs/forward_nakaoEU_weights_{code}.csv')
    okn &= list(a.varId) == list(b_.varId) and np.allclose(a.by, b_.by) and np.allclose(a.bx, b_.bx_codd) and bool(b_.same_sign_as_codd.all())
add('S81: Nakao-weight inputs use the same instruments and outcome rows as the primary analysis, with Codd weights retained for comparison', bool(okn), '4 outcome files')
ST = pd.read_csv('step11_steiger_reverse_removed.csv')
rS_ = R[(R.code == 'SBP->L_codd') & (R.analysis == 'r2<0.001 (primary)') & (R.method == 'IVW (MRE, t)')].iloc[0]; rSs_ = R[(R.code == 'SBP->L_codd') & (R.analysis == 'r2<0.001, Steiger-filtered') & (R.method == 'IVW (MRE, t)')].iloc[0]
s_ = ST[ST.reverse_exposure == 'SBP']
add('S76: Steiger-removed reverse SBP variants account for k and reproduce the Steiger-filtered estimate', len(s_) == int(rS_.nsnp - rSs_.nsnp) and abs(s_.ivw_without_all_removed.iloc[0] - rSs_.b) < 1e-10, f'{len(s_)} variants: ' + ', '.join(f'{r.rsid} ({r.codd_locus})' for r in s_.itertuples()))
CCo = pd.read_csv('step11_conditional_coloc.csv')
add('S82: conditional colocalization completed for every mode x locus x window x trait x LD configuration without error', CCo.note.isna().all() and CCo.groupby(['mode', 'locus', 'window_kb', 'trait', 'ld']).ngroups == 48, f'{len(CCo)} signal pairs; {CCo.groupby(["mode", "locus", "window_kb", "trait", "ld"]).ngroups} configurations')
HI = pd.read_csv('step11_hypertension_se_ivw.csv'); okh = True
for r in HI[HI.outcome_se.str.startswith('P-derived')].itertuples():
    st = R[(R.code == r.code) & (R.analysis == 'r2<0.001 (primary)') & (R.method == 'IVW (MRE, t)')].iloc[0]; okh &= abs(r.b - st.b) < 1e-9 and abs(r.se - st.se) < 1e-9
add('S79: the hypertension SE sensitivity script reproduces the primary IVW results before changing the outcome SE', bool(okh), f'{len(HI)} rows')


# ---- v10 checks (review of v9)
from scipy import stats as _st
okw = True; worst = 0
for r in RB.itertuples():
    if r.method == 'Contamination mixture': continue
    df = (r.k - 1) if 'IVW' in r.method else ((r.k - 2) if 'Egger' in r.method else None)
    p = 2 * (_st.t.sf(abs(r.b / r.se), df) if df else _st.norm.sf(abs(r.b / r.se))); rel = abs(p - r.p) / r.p; worst = max(worst, rel); okw &= rel < 1e-6
add('S80: every non-contamination-mixture P value recomputes from b and SE with its reference distribution', bool(okw), f'{int((RB.method != "Contamination mixture").sum())} rows; max rel diff {worst:.1e}')
CK = pd.read_csv('step12_conmix_check_S85.csv')
okc = (abs(CK.est_own - CK.est_pkg) < 1e-12).all() and (abs(CK.p_package_rule - CK.p_pkg) / CK.p_pkg < 1e-8).all() and (CK[CK.phi == 1].ci_own == CK[CK.phi == 1].ci_pkg).all()
add('S85: contamination-mixture estimate and package P reproduced from an independent profile likelihood (CI exactly where phi = 1)', bool(okc), '; '.join(f'{r.code}: phi {r.phi:.3f}, P pkg {r.p_pkg:.2e}, P (LR/phi^2) {r.p_LR_divided_by_phi2:.2e}' for r in CK.itertuples()))
HC = pd.read_csv('step12_htn_coloc_model_se.csv'); AB_ = pd.read_csv('step6_coloc_v4_abf.csv'); okh2 = True; worst = 0
for r in HC[HC.se_source == 'P-derived (primary)'].itertuples():
    a = AB_[(AB_.locus == r.locus) & (AB_.trait == 'HYPERTENSION') & (AB_.window_kb == r.window_kb) & (np.isclose(AB_.p12, r.p12))].PP4.iloc[0]; worst = max(worst, abs(a - r.PP4)); okh2 &= abs(a - r.PP4) < 1e-3
add('S84: hypertension coloc script reproduces the primary coloc.abf PP4 (P-derived SEs) before changing the SE', bool(okh2), f'{int((HC.se_source == "P-derived (primary)").sum())} runs; max abs diff {worst:.1e} (variants with MAF < 1% dropped)')
WC_ = pd.read_csv('step12_window_check.csv')
add('S83: +/-50 and +/-100 kb windows differ in variant count for every locus and trait', bool((WC_.nsnps_100kb > WC_.nsnps_50kb).all()), f'max |PP4 difference| {WC_.abs_diff_PP4.max():.4f}')
# ---- v13 checks (review of v11)
DG = pd.read_csv('step14_dbgap_mvp_match.csv'); nm = DG.groupby(['portal_dataset', 'portal_trait'])['match'].sum()
add('S86: every MVP portal record (EU, AA, HS; SBP, DBP, hypertension, essential hypertension) matches exactly one dbGaP analysis by N', bool((nm == 1).all() and len(nm) == 12), '; '.join(f'{r.portal_dataset.split("_")[-1]} {r.portal_trait}: {r.dbgap_analysis} ({r.dbgap_accession})' for r in DG[DG['match']].itertuples()))
_bp = DG[DG['match'] & DG.portal_trait.isin(['SBP', 'DBP'])]
add('S86: all matched MVP BP records are the per-participant mean phenotypes', bool(_bp.dbgap_analysis.str.contains('_Mean_INT').all()), ', '.join(_bp.dbgap_analysis))
HQ = pd.read_csv('step14_htn_conditional_coloc.csv'); _dq = HQ.groupby(['mode', 'locus', 'window_kb', 'ld', 'hit1', 'hit2'])['PP.H4'].agg(lambda z: z.max() - z.min()).max()
add('S87: MVP hypertension conditional coloc ran for all 16 settings without error, and P-derived and model SEs give identical posteriors', bool(HQ.note.isna().all() and HQ.groupby(['mode', 'locus', 'window_kb', 'ld']).ngroups == 16 and _dq < 1e-9), f'max |PP4 difference| between SE sources {_dq:.1e}; TERC lead pair PP4 {HQ[(HQ.locus == "TERC") & (HQ.hit1_rsid == "rs2293607") & (HQ.hit2_rsid == "rs12630450")]["PP.H4"].min():.3f}-{HQ[(HQ.locus == "TERC") & (HQ.hit1_rsid == "rs2293607") & (HQ.hit2_rsid == "rs12630450")]["PP.H4"].max():.3f}')
_rf = json.load(open('refs_v4.json'))['nakao26']
add('Reference Nakao 2026 matches PubMed PMID 41896353 (Nat Genet 2026;58(4):831-840; doi:10.1038/s41588-026-02567-1; checked 3 Oct 2026)', '2026;58:831-840' in _rf and '10.1038/s41588-026-02567-1' in _rf, _rf)
_bd = open('manuscript_body_v13.py').read()
CPF = pd.read_csv('step14_conmix_profile.csv').set_index('code'); C85 = pd.read_csv('step12_conmix_check_S85.csv').set_index('code')
add('S88: contamination-mixture profile reproduces S85 (estimate, LR at zero) and the LR P is larger than the CI-implied Wald P for every outcome', bool((abs(CPF.LR_at_0 - C85.loglik_ratio) < 1e-6).all() and (abs(CPF.estimate - C85.est_own) < 1e-9).all() and (CPF.P_LR >= CPF.P_wald_from_CI).all()), '; '.join(f'{c}: LR0 {r.LR_at_0:.1f} vs quadratic {r.LR_quadratic_at_0:.1f}; valid {r.n_valid_at_estimate}->{r.n_valid_at_0}' for c, r in CPF.iterrows()))
add('Main text gives one MVP hypertension effective N (PheCode 401, 318,398); the essential-hypertension record (pha005550) is not cited in the main text', 'essential-hypertension record' not in _bd and 'pha005550' not in _bd, '')
add('Superseded statements removed: "does not give case and control counts" and "case fraction ... not available"', 'does not give case and control counts' not in _bd and 'which its record does not give' not in _bd and 'case fraction not available' not in open('build_supplement_v13.py').read(), '')

def sha(p): return hashlib.sha256(open(p, 'rb').read()).hexdigest()
pats = ['source_data/*', 'ldref/*.json', 'mr_v4_inputs/*.csv', 'mr_v6_inputs/*.csv', 'mr_v9_inputs/*.csv', 'ref_MVMR/*.R', '*.py', '*.R', 'run_all.sh', 'step*.csv', 'presso_equiv/*', 'plink_clump/*.log', 'plink_clump/*.clumped', 'plink_clump/*_kept.txt', 'release_v13/environment.txt', 'step10_render_pdfs.sh', 'step11_render_pdfs.sh', 'step12_render_pdfs.sh', 'step13_render_pdfs.sh', 'step14_render_pdfs.sh', 'step15_render_pdfs.sh', 'INDEPENDENT_RERUN_PROTOCOL.md', 'qhet_bootstrap/*', 'audit_v6_plink/*', 'figures/proof_report.csv']
files = sorted({f for p in pats for f in glob.glob(p) if os.path.isfile(f) and not f.startswith('scratch_')})
render = '--render' in sys.argv
man = pd.DataFrame([dict(path=f, sha256=sha(f), bytes=os.path.getsize(f)) for f in files])
rid = 'v13-' + hashlib.sha256(''.join(man.path + man.sha256).encode()).hexdigest()[:12]
if render:
    rid = json.load(open('release_v13/release_summary.json'))['release_id']
    import docx
    txt = '\n'.join(p.text for p in docx.Document('out/Manuscript_TERC_TERT_BP_v13_2026-10-03.docx').paragraphs)
    for t in docx.Document('out/Manuscript_TERC_TERT_BP_v13_2026-10-03.docx').tables: txt += '\n'.join(c.text for row in t.rows for c in row.cells)
    g_ = lambda code, an, m: R[(R.code == code) & (R.analysis == an) & (R.method == m)].iloc[0]
    want = []
    for code in ['S_mvpEU', 'D_mvpEU']:
        r = g_(code, 'r2<0.001 (primary)', 'IVW (MRE, t)'); want.append(f'{r.b:.3f}'); want.append(f'{r.lo:.3f} to {r.hi:.3f}')
    for code in ['SBP->L_codd', 'DBP->L_codd', 'Hypertension->L_codd']:
        r = g_(code, 'r2<0.001 (primary)', 'IVW (MRE, t)'); want.append(f'{r.b:.4f}'.replace('-', '−'))
    r = g_('H_mvpEU', 'r2<0.001 (primary)', 'IVW (MRE, t)'); want.append(f'OR {np.exp(r.b):.2f}')
    mv = E21[(E21.outcome == 'S_mvpEU') & (E21.exposure == 'LTL') & (E21.method == 'MV-IVW (random, t)')].iloc[0]; want.append(f'{mv.b:.3f}')
    for oc in ['S_mvpEU']:
        q = Q[(Q.outcome == oc) & (Q.exposure == 'LTL') & (Q.rho == 0)].iloc[0]; want.append(f'{q.pct_lo:.3f} to {q.pct_hi:.3f}')
    for code in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU']:
        want.append(f'{int(CR[(CR.code == code) & (CR.B == 1000000)].n_outliers.iloc[0])}')
    want += [f'{RB[(RB.code == "S_mvpEU") & (RB.method == m)].b.iloc[0]:.3f}' for m in ['Contamination mixture', 'cML-MA-BIC-DP', 'MR-RAPS (Huber loss, over-dispersion)']]
    want += [f'{RB[(RB.code == "S_mvpEU") & (RB.weights != "Codd 2021") & (RB.method == "IVW (MRE, t)")].b.iloc[0]:.3f}', f'{rSs_.b:.4f}'.replace('-', '−'), 'rs12638862', 'rs555185638', 'rs3891167', 'rs7712562']
    miss = [w for w in want if w not in txt.replace('-', '−') and w not in txt]
    add('Rendered manuscript contains the key values of the result files (primary, reverse, MVMR, qhet, MR-PRESSO, robust estimators, Nakao weights, Steiger variants)', not miss, f'{len(want) - len(miss)}/{len(want)} values found' + (f'; missing {miss}' if miss else ''))
    stale_txt = [s for s in ['validated re-implementation', 'prespecified stability', 'explains the TERC', 'all 330 variants', 'Keaton2024 2024', 'v5', 'v6 ', 'reproduced MRPRESSO 1.0 exactly', 'remove Monte Carlo instability', 'MR exposure (all forward', 'with unresolved multi-signal sensitivity; regional', 'stable at ±100 kb only', 'Secondary FinnGen analysis:', 'in the record)', 'More simulations', 'This apparent contrast', 'Genetically predicted liability differs', 'colocalization at TERC and TERT', '-0.0072', 'whereas 30 do not', 'We want to acknowledge', 'could not be retrieved', 'per unit of genetically instrumented', 'whose exposure unit was not stated', 'which would carry the forward'] if s in txt]
    add('Rendered manuscript contains no withdrawn wording or version labels', not stale_txt, ', '.join(stale_txt) or 'none')
    nest = [m.group(0) for m in re.finditer(r'\([^()]*\(95% CI[^()]*\)', txt)] + [m.group(0) for m in re.finditer(r'\) \((?:P|k) = ', txt)]
    add('Rendered manuscript: no estimate CI nested inside parentheses and no doubled parentheses before P or k', not nest, f'{len(nest)} found' + (f': {nest[:3]}' if nest else ''))
    _nref = len(re.findall(r'^\d+\. ', txt.split('\nReferences\n')[1].split('\nTables')[0], flags=re.M)) if '\nReferences\n' in txt else -1
    add('Rendered manuscript: main-text references <= 50 (Circ Genom Precis Med limit)', 0 < _nref <= 50, f'{_nref} references')
    _body = txt.split('\nReferences\n')[0]; _seen = []
    for _m in re.finditer(r'\[(\d+(?:[,–-]\d+)*)\]', _body):
        for _part in _m.group(1).split(','):
            _ab = [int(x) for x in re.split('[–-]', _part)]
            for _n in range(_ab[0], _ab[-1] + 1):
                if _n not in _seen: _seen.append(_n)
    add('Main-text citations are numbered in order of first appearance (1, 2, 3, ...)', _seen == list(range(1, len(_seen) + 1)), f'{len(_seen)} distinct citations in the text')
    add('STROBE-MR reference carries the full published title', 'Using Mendelian Randomization: The STROBE-MR Statement. JAMA. 2021;326:1614-1621' in txt, '')
    _ab = txt[txt.index('Background.'):txt.index('Keywords')]
    add('Rendered manuscript: abstract <= 250 words', len(_ab.split()) <= 250, f'{len(_ab.split())} words')
    add('Rendered manuscript: Methods include TOP Guidelines and IRB statements', ('Transparency and Openness Promotion' in txt) and ('Institutional Review Board' in txt), 'present' if ('Transparency and Openness Promotion' in txt) else 'missing')
    import zipfile
    fig_h = {sha(f): f for f in glob.glob('figures/Fig*.png')}
    det = []; oke = True
    def _all(f):
        _d = docx.Document(f); _t = '\n'.join(p.text for p in _d.paragraphs)
        for _tb in _d.tables: _t += '\n' + '\n'.join(c.text for r_ in _tb.rows for c in r_.cells)
        return _t
    _tm, _ts = _all('out/Manuscript_TERC_TERT_BP_v13_2026-10-03.docx'), _all('out/Supplementary_Information_v13_2026-10-03.docx')
    add('Manuscript and Supplementary Information come from one build: both carry this release ID', rid in _ts and rid in _tm, rid)
    _brit = re.compile(r'\b(analysed|analyse|modelled|modelling|labelled|neighbouring|colour\w*|minimis\w*|favour\w*|behaviour\w*|haemat\w*|harmonis\w*|summaris\w*|vectoris\w*|remodelling)\b', re.I)
    _hits = sorted(set(m.group(0) for m in _brit.finditer(_tm + _ts)))
    add('US spelling in manuscript and Supplementary Information text and tables (reference titles excepted)', not _hits, ', '.join(_hits))
    add('Supplementary Methods heading no longer labelled by review round', 'Review-response checks (v10)' not in _ts, '')
    _cit = set()
    for _m in re.finditer(r'Supplementary Tables? ((?:S\d+(?:[–-]S\d+)?(?:, | and |,? and )?)+)', _tm + _ts):
        for _a, _b in re.findall(r'S(\d+)(?:[–-]S(\d+))?', _m.group(1)):
            _a = int(_a); _b = int(_b or _a)
            if _b - _a <= 20: _cit.update(range(_a, _b + 1))
    _miss = sorted(set(range(1, 89)) - _cit)
    add('Every Supplementary Table S1-S88 is cited individually (not only by the S1-S88 range)', not _miss, ', '.join(f'S{x}' for x in _miss))
    _cph = pd.read_csv('step12_conmix_check_S85.csv').set_index('code')
    _pd = f'{_cph.loc["D_mvpEU"].p_LR_divided_by_phi2:.1e}'.split('e')
    _dbp_ok = ('package P' in _tm) and (f'{float(_pd[0]):.1f} × 10{int(_pd[1])}'.replace('-', '−') in _tm or f'{_cph.loc["D_mvpEU"].p_LR_divided_by_phi2:.3f}' in _tm)
    add('Table 2 reports the DBP contamination-mixture P with phi^2 applied as for the CI, package value in Notes', bool(_dbp_ok), f'corrected P {_cph.loc["D_mvpEU"].p_LR_divided_by_phi2:.2e}; package {_cph.loc["D_mvpEU"].p_package_rule:.2e}')
    _dm = docx.Document('out/Manuscript_TERC_TERT_BP_v13_2026-10-03.docx')
    _ln = all(sec._sectPr.find(qn_('w:lnNumType')) is not None for sec in _dm.sections); _pg = all('PAGE' in sec.footer._element.xml for sec in _dm.sections)
    add('Manuscript has continuous line numbers and page numbers in every section', _ln and _pg, f'{len(_dm.sections)} sections')
    import subprocess as _sp
    _meta = {f: _sp.run(['pdfinfo', f], capture_output=True, text=True).stdout for f in sorted(glob.glob('out/*v13*.pdf'))}
    _bad = [os.path.basename(f) for f, m in _meta.items() if 'python-docx' in m or not re.search(r'^Title:\s+\S', m, re.M)]
    add('PDF metadata: every PDF has a title and an author other than python-docx', not _bad and len(_meta) == 5, ', '.join(_bad) or f'{len(_meta)} PDFs')
    _wb = pd.read_excel('out/Supplementary_Tables_TERC_TERT_BP_v13.xlsx', sheet_name='S73_build_checks')
    add('Workbook S73 sheet contains no FAIL and has the same release ID', bool((_wb.status == 'PASS').all() and (_wb.release_id == rid).all()), f'{int((_wb.status == "PASS").sum())}/{len(_wb)} PASS in S73')
    _pdf = _sp.run(['pdftotext', 'out/Manuscript_TERC_TERT_BP_v13_2026-10-03.pdf', '-'], capture_output=True, text=True).stdout.split('\f')[:-1]
    _imgp = set(int(l.split()[0]) for l in _sp.run(['pdfimages', '-list', 'out/Manuscript_TERC_TERT_BP_v13_2026-10-03.pdf'], capture_output=True, text=True).stdout.splitlines()[2:] if l.split())
    _blank = [i + 1 for i, pg in enumerate(_pdf) if not re.sub(r'\s|Page \d+|\d+', '', pg) and (i + 1) not in _imgp]
    add('Manuscript PDF has no blank pages (no text other than page/line numbers and no image)', not _blank, ', '.join(map(str, _blank)) or f'{len(_pdf)} pages')
    for docf in ['out/Manuscript_TERC_TERT_BP_v13_2026-10-03.docx', 'out/Supplementary_Information_v13_2026-10-03.docx']:
        z = zipfile.ZipFile(docf); media = [n for n in z.namelist() if n.startswith('word/media/')]
        hs = [hashlib.sha256(z.read(n)).hexdigest() for n in media]; found = [fig_h.get(h) for h in hs]
        oke &= all(found) and len(media) > 0; det.append(f'{os.path.basename(docf)}: {sum(1 for x in found if x)}/{len(media)} images identical to release figures')
    add('Figures embedded in the manuscript and Supplementary Information are byte-identical to the release figure files', oke, '; '.join(det))
    PR_ = pd.read_csv('figures/proof_report.csv'); main_ = PR_[PR_.figure.str.match(r'^Fig\d')]
    add('Main figures at 180 mm print width: no text below 6 pt', len(main_) == 4 and (main_.min_effective_pt.astype(float) >= 6.0 - 1e-9).all(), '; '.join(f'{r.figure}: min {float(r.min_effective_pt):.2f} pt, scale {float(r.scale_at_180mm):.3f}' for r in main_.itertuples()))
    outs = sorted(glob.glob('out/*v13*') + glob.glob('figures/Fig*.p*'))
    man = pd.concat([man, pd.DataFrame([dict(path=f, sha256=sha(f), bytes=os.path.getsize(f)) for f in outs])], ignore_index=True)
    prev = pd.read_csv('release_v13/build_checks.csv'); prev = prev[prev.phase == 'pre-build']
    CC = pd.concat([prev, pd.DataFrame(C).assign(release_id=rid, phase='post-build')], ignore_index=True)
else:
    CC = pd.DataFrame(C).assign(release_id=rid, phase='pre-build')
CC = CC[['release_id', 'phase', 'check', 'status', 'detail']]; CC.to_csv('release_v13/build_checks.csv', index=False); man.to_csv('release_v13/release_manifest.csv', index=False)
pre = CC[CC.phase == 'pre-build']
json.dump(dict(release_id=rid, n_checks=int(len(pre)), n_checks_pass=int((pre.status == 'PASS').sum()), n_post=int((CC.phase == 'post-build').sum()),
               n_post_pass=int(((CC.phase == 'post-build') & (CC.status == 'PASS')).sum()), n_files=int(len(man))), open('release_v13/release_summary.json', 'w'), indent=1)
print(CC.to_string()); print('release', rid, 'files', len(man))
if (CC.status != 'PASS').any(): sys.exit('RELEASE CHECKS FAILED')
