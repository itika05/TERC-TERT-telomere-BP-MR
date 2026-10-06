"""v7 release checks and manifest. Run after all analyses (before the documents are built): machine-readable PASS/FAIL checks that the
publication tables belong to one run and reconstruct from their named inputs, plus a SHA-256 manifest and a release ID derived from the hashes
of all inputs, scripts and result files. With --render (after the documents are built) it adds checks that key values in the rendered
manuscript equal the result files, and hashes the rendered outputs. The build fails if any check fails."""
import sys, os, glob, json, hashlib, re
import numpy as np, pandas as pd
from scipy import stats
os.makedirs('release_v7', exist_ok=True)
C = []
def add(check, ok, detail=''): C.append(dict(check=check, status='PASS' if ok else 'FAIL', detail=detail))
def ivw(bx, by, sy):
    w = 1 / sy ** 2; b = (w * bx * by).sum() / (w * bx ** 2).sum(); k = len(bx)
    rss = (w * (by - b * bx) ** 2).sum(); se = np.sqrt(max(1.0, rss / (k - 1)) / (w * bx ** 2).sum()); p = 2 * stats.t.sf(abs(b / se), k - 1)
    return b, se, p, rss, k
R = pd.read_csv('step8_mr_results_v6.csv')
key = R.direction.astype(str) + '|' + R.code + '|' + R.analysis + '|' + R.method + '|' + R.source.astype(str)
add('S11 analysis keys unique (direction|code|analysis|method|source)', not key.duplicated().any(), f'{len(R)} rows')
O = pd.read_csv('step9_mr_output_consistency_S46.csv'); add('S46: every S11 row reproduces P and CI from b and SE', (O.status == 'PASS').all() and len(O) == len(R), f'{(O.status == "PASS").sum()}/{len(O)} rows')
T = pd.read_csv('step9_reverse_reconstruction_test.csv'); add('S72: every reverse IVW result reconstructs from S13 rows (k, b, SE, P)', (T.status == 'PASS').all(), f'{len(T)} results; max rel diff {T.max_rel_diff.max():.1e}')
# forward IVW from S12 rows (primary, FinnGen) and PLINK forward inputs
S12 = pd.read_csv('step6_v4_harmonised_inputs_all.csv'); worst = 0; n = 0; okk = True
for code in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU', 'H_fg12']:
    d = S12[(S12.code == code) & (S12.file == f'forward_selected_r2_0.001_{code}.csv')]
    b, se, p, rss, k = ivw(d.bx.values, d.by.values, d.sy.values); st = R[(R.code == code) & (R.analysis == 'r2<0.001 (primary)') & (R.method == 'IVW (MRE, t)')].iloc[0]
    rel = max(abs(b - st.b) / abs(st.b), abs(se - st.se) / st.se); worst = max(worst, rel); n += 1; okk &= (k == st.nsnp) and rel < 1e-8
    d = pd.read_csv(f'mr_v6_inputs/forward_plink_{code}.csv'); b, se, p, rss, k = ivw(d.bx.values, d.by.values, d.sy.values)
    st = R[(R.code == code) & (R.analysis == 'PLINK 1.9 clumping') & (R.method == 'IVW (MRE, t)')].iloc[0]
    rel = max(abs(b - st.b) / abs(st.b), abs(se - st.se) / st.se); worst = max(worst, rel); n += 1; okk &= (k == st.nsnp) and rel < 1e-8
add('Forward IVW (primary from S12 rows; PLINK sets) reconstructs S11', okk, f'{n} results; max rel diff {worst:.1e}')
# MR-PRESSO corrected estimates at B = 1e6 reconstruct from S12 minus the listed outliers
CR = pd.read_csv('step8_presso_convergence_runs.csv'); okp = True; det = []
for code in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU']:
    r6 = CR[(CR.code == code) & (CR.B == 1000000)]; sets = r6.outliers.fillna('').map(lambda x: frozenset(x.split(';')) - {''})
    same = sets.nunique() == 1; okp &= same
    d = S12[(S12.code == code) & (S12.file == f'forward_selected_r2_0.001_{code}.csv')]; keep = ~d.rsid.isin(sets.iloc[0])
    dd = d[keep]; w = 1 / dd.sy ** 2; b = (w * dd.bx * dd.by).sum() / (w * dd.bx ** 2).sum()
    st = R[(R.code == code) & (R.method == 'MR-PRESSO (outlier-corrected)')].iloc[0]; okp &= abs(b - st.b) < 1e-10 and int(keep.sum()) == int(st.nsnp)
    det.append(f'{code}: {len(sets.iloc[0])} outliers, k {int(keep.sum())}')
add('MR-PRESSO B = 1e6: identical outlier sets over 5 seeds; corrected b reconstructs from S12', okp, '; '.join(det))
E = pd.read_csv('step9_presso_equivalence.csv'); E = E[E.check != 'summary']
add('S70: Python MR-PRESSO equals MRPRESSO 1.0 on shared simulated data (all intermediates)', len(E) > 0 and (E.status == 'PASS').all() and E.code.nunique() == 3, f'{(E.status == "PASS").sum()}/{len(E)} checks, {E.code.nunique()} outcomes')
# MVMR: weighted regression on S20 minus S39 reproduces S21 (all three exposures)
M = pd.read_csv('mr_v4_inputs/mvmr_instruments.csv'); fb = set(pd.read_csv('step7_mvmr_fallback_variants.csv').varId); M = M[~M.varId.isin(fb)]
E21 = pd.read_csv('step9_mvmr_estimates_primary.csv'); okm = len(M) == 318; worst = 0
for oc in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU']:
    X = M[['LTL_b', 'BMI_b', 'LYM_b']].values; y = M[oc + '_b'].values; w = 1 / M[oc + '_se'].values ** 2
    beta = np.linalg.solve(X.T @ (X * w[:, None]), X.T @ (w * y))
    for j, ex in enumerate(['LTL', 'BMI', 'LYM']):
        st = E21[(E21.outcome == oc) & (E21.exposure == ex) & (E21.method == 'MV-IVW (random, t)')].iloc[0]
        rel = abs(beta[j] - st.b) / abs(st.b); worst = max(worst, rel); okm &= rel < 1e-8 and int(st.nsnp) == 318
add('S21: MV-IVW coefficients (LTL, BMI, LYM) reconstruct from S20 minus S39 (318 variants)', okm, f'max rel diff {worst:.1e}')
Q = pd.read_csv('step9_mvmr_qhet.csv'); g = Q.groupby(['outcome', 'rho'])
okq = (g.size() == 3).all() and (g.exposure.apply(lambda s: list(s) == ['LTL', 'BMI', 'LYM'])).all() and (g.qhet_estimate.nunique() == 3).all()
V = pd.read_csv('step9_qhet_vectorised_check.csv'); okq &= V.abs_diff.max() < 1e-10
add('S55: qhet_mvmr exported one row per exposure (LTL, BMI, LYM) per outcome and rho; three distinct values; vectorised objective equals reference', bool(okq), f'{len(g)} outcome x rho groups')
# variant accounting for reverse PLINK sets before/after outcome matching
RL = pd.read_csv('step8_reverse_source_ledger.csv'); okv = True; det = []
for r in RL.itertuples():
    an = 'r2<0.001 (primary)' if r.selection == 'custom' else 'PLINK 1.9 clumping'
    st = R[(R.code == f'{r.exposure}->L_codd') & (R.analysis == an) & (R.method == 'IVW (MRE, t)')].iloc[0]
    okv &= (r.instruments - r.absent - r.allele_mismatch) == st.nsnp; det.append(f'{r.exposure}/{r.selection}: {r.instruments} selected, {r.absent} without source record, k {int(st.nsnp)}')
add('Reverse variant accounting: selected - unmatched = k used', okv, '; '.join(det))
L52 = pd.read_csv('step8_plink_discrepancy_ledger.csv')
okl = L52.plink_r2_dprime_ML.notna().all() and (L52.mechanism.str.startswith('threshold') | L52.mechanism.str.startswith('selection chain')).all()
c = L52[L52.plink_clump_r2.notna()]; okl &= ((c.plink_clump_r2 - c.plink_r2_dprime_ML).abs() / c.plink_clump_r2).max() < 0.005
add('S52: every discrepant variant has a PLINK --r2 dprime value and an explained mechanism; dprime r2 matches clump r2 (<0.5%)', bool(okl), f'{len(L52)} rows')
L38 = pd.read_csv('step9_se_fallback_ledger_S38.csv')
add('S38: no reverse fallback record is marked as used; MVMR fallback variants all excluded', (L38[L38.direction == 'reverse'].status_current_release.str.startswith('superseded')).all() and set(L38[L38.direction == 'MVMR'].varId).isdisjoint(set(M.varId)), f'{len(L38)} records')
# no stale sources in the publication workbook builder
bs = open('build_supplement_v7.py').read()
stale = [f for f in ['step6_v4_reverse_inputs_all.csv', 'step7_se_fallback_ledger.csv', 'step7_mr_output_consistency.csv', 'step8_mvmr_v6_qhet.csv', 'step8_mvmr_v6_estimates.csv'] if f"'{f}'" in bs]
add('Workbook uses no superseded source for a current table', not stale, ', '.join(stale) or 'none')

def sha(p): return hashlib.sha256(open(p, 'rb').read()).hexdigest()
pats = ['source_data/*', 'ldref/*.json', 'mr_v4_inputs/*.csv', 'mr_v6_inputs/*.csv', 'ref_MVMR/*.R', '*.py', '*.R', 'run_all.sh', 'step*.csv', 'presso_equiv/*', 'plink_clump/*.log', 'plink_clump/*.clumped', 'plink_clump/*_kept.txt', 'release_v7/environment.txt']
files = sorted({f for p in pats for f in glob.glob(p) if os.path.isfile(f) and not f.startswith('scratch_')})
render = '--render' in sys.argv
man = pd.DataFrame([dict(path=f, sha256=sha(f), bytes=os.path.getsize(f)) for f in files])
rid = 'v7-' + hashlib.sha256(''.join(man.path + man.sha256).encode()).hexdigest()[:12]
if render:
    rid = json.load(open('release_v7/release_summary.json'))['release_id']
    import docx
    txt = '\n'.join(p.text for p in docx.Document('out/Manuscript_TERC_TERT_BP_v7_2026-09-30.docx').paragraphs)
    for t in docx.Document('out/Manuscript_TERC_TERT_BP_v7_2026-09-30.docx').tables: txt += '\n'.join(c.text for row in t.rows for c in row.cells)
    g_ = lambda code, an, m: R[(R.code == code) & (R.analysis == an) & (R.method == m)].iloc[0]
    want = []
    for code in ['S_mvpEU', 'D_mvpEU']:
        r = g_(code, 'r2<0.001 (primary)', 'IVW (MRE, t)'); want.append(f'{r.b:.3f}'); want.append(f'{r.lo:.3f} to {r.hi:.3f}')
    for code in ['SBP->L_codd', 'DBP->L_codd', 'Hypertension->L_codd']:
        r = g_(code, 'r2<0.001 (primary)', 'IVW (MRE, t)'); want.append(f'{r.b:.4f}'.replace('-', '−'))
    r = g_('H_mvpEU', 'r2<0.001 (primary)', 'IVW (MRE, t)'); want.append(f'OR {np.exp(r.b):.2f}')
    mv = E21[(E21.outcome == 'S_mvpEU') & (E21.exposure == 'LTL') & (E21.method == 'MV-IVW (random, t)')].iloc[0]; want.append(f'{mv.b:.3f}')
    for oc in ['S_mvpEU']:
        q = Q[(Q.outcome == oc) & (Q.exposure == 'LTL') & (Q.rho == 0)].iloc[0]; want.append(f'{q.pct_lo:.3f} to {q.pct_hi:.3f}')
    for code in ['S_mvpEU', 'D_mvpEU', 'H_mvpEU']:
        want.append(f'{int(CR[(CR.code == code) & (CR.B == 1000000)].n_outliers.iloc[0])}')
    miss = [w for w in want if w not in txt.replace('-', '−') and w not in txt]
    add('Rendered manuscript contains the key values of the result files (primary, reverse, MVMR, qhet, MR-PRESSO)', not miss, f'{len(want) - len(miss)}/{len(want)} values found' + (f'; missing {miss}' if miss else ''))
    stale_txt = [s for s in ['validated re-implementation', 'prespecified stability', 'explains the TERC', 'all 330 variants', 'Keaton2024 2024', 'v5', 'v6 '] if s in txt]
    add('Rendered manuscript contains no withdrawn wording or version labels', not stale_txt, ', '.join(stale_txt) or 'none')
    outs = sorted(glob.glob('out/*v7*') + glob.glob('figures/Fig*.p*'))
    man = pd.concat([man, pd.DataFrame([dict(path=f, sha256=sha(f), bytes=os.path.getsize(f)) for f in outs])], ignore_index=True)
    prev = pd.read_csv('release_v7/build_checks.csv'); prev = prev[prev.phase == 'pre-build']
    CC = pd.concat([prev, pd.DataFrame(C).assign(release_id=rid, phase='post-build')], ignore_index=True)
else:
    CC = pd.DataFrame(C).assign(release_id=rid, phase='pre-build')
CC = CC[['release_id', 'phase', 'check', 'status', 'detail']]; CC.to_csv('release_v7/build_checks.csv', index=False); man.to_csv('release_v7/release_manifest.csv', index=False)
pre = CC[CC.phase == 'pre-build']
json.dump(dict(release_id=rid, n_checks=int(len(pre)), n_checks_pass=int((pre.status == 'PASS').sum()), n_post=int((CC.phase == 'post-build').sum()),
               n_post_pass=int(((CC.phase == 'post-build') & (CC.status == 'PASS')).sum()), n_files=int(len(man))), open('release_v7/release_summary.json', 'w'), indent=1)
print(CC.to_string()); print('release', rid, 'files', len(man))
if (CC.status != 'PASS').any(): sys.exit('RELEASE CHECKS FAILED')
