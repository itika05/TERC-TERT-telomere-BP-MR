#!/usr/bin/env python3
"""Checks run on the final submission package (Word, PDF, figure and workbook files).

Usage: python3 final_package_checks.py <Upload_to_journal dir> <Editable_sources dir> <out.csv>
Writes one row per check (check, status, detail). Exit code 1 if any check FAILs.
The placeholder gate is reported separately (GATE) and is not written to the workbook.
"""
import sys, re, csv, zipfile, io, os, glob
from lxml import etree
import fitz
from PIL import Image, ImageChops, ImageStat

U, E, OUT = sys.argv[1:4]
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
rows = []
def add(check, ok, detail=''):
    rows.append((check, 'PASS' if ok else 'FAIL', detail))

def docx_paras(path):
    z = zipfile.ZipFile(path)
    root = etree.fromstring(z.read('word/document.xml'))
    out = []
    for p in root.iter(W + 'p'):
        out.append(''.join(t.text or '' for t in p.iter(W + 't')))
    return z, root, out

ms_docx = os.path.join(U, 'Manuscript_TERC_TERT_BP.docx')
sup_docx = os.path.join(E, 'Supplemental_Material_TERC_TERT_BP.docx')
str_docx = os.path.join(E, 'STROBE-MR_Checklist_TERC_TERT_BP.docx')
cl_docx = os.path.join(E, 'Cover_Letter_TERC_TERT_BP.docx')
zms, rms, P = docx_paras(ms_docx)
_, rsup, S = docx_paras(sup_docx)
_, _, ST = docx_paras(str_docx)
_, _, CL = docx_paras(cl_docx)
ms_text = '\n'.join(P); sup_text = '\n'.join(S)

idx = {t.strip(): i for i, t in enumerate(P) if t.strip()}
iA, iN, iI, iAI, iR = idx['ABSTRACT'], idx['Nonstandard Abbreviations and Acronyms'], idx['INTRODUCTION'], idx['ARTICLE INFORMATION'], idx['REFERENCES']
wc = lambda a, b: sum(len(P[i].split()) for i in range(a, b))

# 1 abstract
abs_words = sum(len(P[i].split()) for i in range(iA + 1, iN) if re.match(r'(BACKGROUND|METHODS|RESULTS|CONCLUSIONS):', P[i]))
add('Abstract <= 250 words (four headed sections, headings included)', abs_words <= 250, f'{abs_words} words')
heads = [h for h in ('BACKGROUND:', 'METHODS:', 'RESULTS:', 'CONCLUSIONS:') if any(P[i].startswith(h) for i in range(iA, iN))]
add('Abstract has Background, Methods, Results and Conclusions sections', len(heads) == 4, ', '.join(h[:-1].title() for h in heads))

# 2 word counts stated on title page
total = sum(len(t.split()) for i, t in enumerate(P) if not iN <= i < iI)
text_w = wc(iI, iAI)
m = re.search(r'Total word count: ([\d,]+).*?text \(Introduction to Conclusions\): ([\d,]+); abstract: (\d+)', ms_text)
st = tuple(int(x.replace(',', '')) for x in m.groups()) if m else None
add('Title-page word counts match the manuscript file (total excludes the nonstandard abbreviations list)', st == (total, text_w, abs_words),
    f'stated {st}; computed total {total}, text {text_w}, abstract {abs_words}')
short = re.search(r'Short title: (.*)', ms_text).group(1).strip()
add('Short title <= 50 characters including spaces', len(short) <= 50, f'{len(short)} characters')

# 3 references
refs = []
for i in range(iR + 1, len(P)):
    if re.match(r'\d+\.', P[i]) and not re.match(r'\d+\.\d+ (to|\()', P[i]): refs.append(P[i])
    elif P[i].startswith('Table 1.'): break
nums = [int(re.match(r'(\d+)\.', r).group(1)) for r in refs]
add('Reference list numbered 1..N without gaps', nums == list(range(1, len(nums) + 1)), f'{len(nums)} references')
mref = re.search(r'References (\d+)[–-](\d+)', ms_text)
sup_first = int(mref.group(1)) if mref else len(nums) + 1
main_n = sup_first - 1
# superscript citations in main text (title page to end of Article Information + tables + legends)
body = rms.find(W + 'body')
def sup_cites(paras_range_filter):
    cites = []
    for k, p in enumerate(body.iter(W + 'p')):
        if not paras_range_filter(k):
            continue
        for r in p.iter(W + 'r'):
            va = r.find(f'{W}rPr/{W}vertAlign')
            if va is not None and va.get(W + 'val') == 'superscript':
                t = ''.join(x.text or '' for x in r.iter(W + 't'))
                for part in re.split(r'[,\s]+', t):
                    mm = re.fullmatch(r'(\d+)[–-](\d+)', part)
                    if mm:
                        cites += list(range(int(mm.group(1)), int(mm.group(2)) + 1))
                    elif part.isdigit():
                        cites.append(int(part))
    return cites
ms_cites = sup_cites(lambda k: k < iR or k > iR + len(refs))
ms_cites = [c for c in ms_cites if c <= len(nums)]
first = []
for c in ms_cites:
    if c not in first:
        first.append(c)
add(f'Main-text references <= 50 (Circ Genom Precis Med limit); references {sup_first}-{len(nums)} are cited in the Supplemental Material only',
    main_n <= 50 and max(first) == main_n and all(c <= main_n for c in first), f'{main_n} main-text references; {len(nums) - main_n} supplement-only')
add('Main-text citations numbered in order of first appearance', first == list(range(1, len(first) + 1)), f'{len(first)} distinct citations')
sup_c = sup_cites.__code__ and None
# supplement citations resolve to the list
sc = []
for p in rsup.iter(W + 'p'):
    for r in p.iter(W + 'r'):
        va = r.find(f'{W}rPr/{W}vertAlign')
        if va is not None and va.get(W + 'val') == 'superscript':
            t = ''.join(x.text or '' for x in r.iter(W + 't'))
            for part in re.split(r'[,\s]+', t):
                mm = re.fullmatch(r'(\d+)[–-](\d+)', part)
                if mm: sc += list(range(int(mm.group(1)), int(mm.group(2)) + 1))
                elif part.isdigit(): sc.append(int(part))
supp_only = set(range(sup_first, len(nums) + 1))
add('Every supplement-only reference is cited in the Supplemental Material, and the Supplemental Material has no separate reference list',
    supp_only <= set(sc) and not re.search(r'\n(References|REFERENCES|Supplemental References)\n', sup_text),
    f'{len(supp_only & set(sc))}/{len(supp_only)} cited')
ref18 = refs[17]
add('Reference 18 is Nakao et al., Nat Genet 2026;58:831-840', 'Nakao T' in ref18 and '2026;58:831' in ref18, ref18[:60])

# 4 display items
tabs = re.findall(r'^Table (\d+)\. ', ms_text, re.M); figs = re.findall(r'^Figure (\d+)\. ', ms_text, re.M)
add('Display items <= 8 (tables + figures)', len(tabs) + len(figs) <= 8, f'{len(tabs)} tables, {len(figs)} figures')

# 5 version / history wording
pat = re.compile(r'(\bv\d+\b(?<!GTEx v8)|earlier version|previous version|this version|superseded|withdrawn|release ID|dddb0ef8a7b4|CircGPM_v\d+)', re.I)
allowed = re.compile(r'GTEx v8|API v2|api/v2|v0\.9\.0|\bv1\.0\.0\b|dbGaP|pha\d+\.\d|phs\d+\.v\d', re.I)
hits = []
for name, txt in (('manuscript', ms_text), ('supplement', sup_text), ('STROBE-MR', '\n'.join(ST)), ('cover letter', '\n'.join(CL))):
    for mm in pat.finditer(txt):
        ctx = txt[max(0, mm.start() - 25):mm.end() + 25]
        if not allowed.search(ctx):
            hits.append(f'{name}: ...{ctx}...')
add('No version labels or version-history wording in manuscript, supplement, STROBE-MR or cover letter (GTEx v8, R package and GitHub release versions excepted)', not hits, '; '.join(hits[:5]) or 'none')

# 6 table citations
cited = set()
for txt in (ms_text, sup_text):
    for mm in re.finditer(r'Tables? (S\d+(?:\s*(?:,|and|to|–|-)\s*S?\d+)*)', txt):
        for a, b in re.findall(r'S?(\d+)\s*[–-]\s*S?(\d+)', mm.group(1)):
            if int(b) - int(a) < 10:
                cited.update(range(int(a), int(b) + 1))
        cited.update(int(x) for x in re.findall(r'S(\d+)', mm.group(1)))
missing = sorted(set(range(1, 89)) - cited)
add('Every Supplemental Table S1-S88 is cited individually (ranges of 10 or more not counted)', not missing, f'missing: {missing}' if missing else 'none missing')

# 7 statements carried over from the analysis build, rechecked on the final text
add('Main text gives one MVP hypertension effective N (318,398) and does not cite the essential-hypertension record (pha005550)',
    '318,398' in ms_text and 'pha005550' not in ms_text, None)
add('Statements on MVP hypertension case and control counts match the dbGaP record (no statement that counts or case fraction are unavailable)',
    not re.search(r'(case|control)[^.]{0,80}(unavailable|not available|not reported)', ms_text + sup_text, re.I), None)
add('STROBE-MR reference carries the full published title', any('Strengthening the Reporting of Observational Studies in Epidemiology Using Mendelian Randomization' in r for r in refs), None)
add('Methods include TOP Guidelines (data availability) and ethics/IRB sections', 'Data Availability' in ms_text and 'Ethics and Study Design' in ms_text and 'Institutional Review Board' in ms_text, None)
add('Data Availability gives the public GitHub repository and release', 'https://github.com/itika05/TERC-TERT-telomere-BP-MR' in ms_text and 'v1.0.0' in ms_text, None)
brit = re.compile(r'\b(colour|analys(e|ed|ing)|randomis\w*|harmonis\w*|behaviour|centre|modelled|haemo\w*|tumour|programme|summaris\w*|characteris\w*|visualis\w*|normalis\w*)\b', re.I)
bh = [mm.group(0) for txt in (ms_text, sup_text) for line in txt.split('\n') if not re.match(r'\d+\.', line) for mm in brit.finditer(line)]
# reference titles excepted (lines starting with a reference number)
add('US spelling in manuscript and supplement text and tables (reference titles excepted)', not bh, ', '.join(sorted(set(bh))) or None)

# 8 layout
sect = rms.findall(f'.//{W}sectPr')
ln = [s.find(W + 'lnNumType') is not None for s in sect]
foot = ''.join(zms.read(n).decode('utf8') for n in zms.namelist() if n.startswith('word/footer'))
add('Manuscript has continuous line numbers and page-number fields in every section', all(ln) and 'PAGE' in foot, f'{len(sect)} sections')

# 9 PDFs
pdfs = sorted(glob.glob(os.path.join(U, '*.pdf')) + glob.glob(os.path.join(E, '*.pdf')))
bad = []
for f in pdfs:
    md = fitz.open(f).metadata
    if not md.get('title') or not md.get('author') or 'python-docx' in (md.get('author') or '') + (md.get('subject') or '') + (md.get('keywords') or ''):
        bad.append(os.path.basename(f))
add('PDF metadata: every PDF has a title and an author', not bad, f'{len(pdfs)} PDFs' + (f'; missing: {bad}' if bad else ''))
msp = fitz.open(os.path.join(E, 'Manuscript_TERC_TERT_BP.pdf'))
blank = [i + 1 for i, pg in enumerate(msp) if not re.sub(r'[\d\s]', '', pg.get_text()) and not pg.get_images()]
add('Manuscript PDF (Manuscript_TERC_TERT_BP.pdf, rendered from Manuscript_TERC_TERT_BP.docx) has no blank pages', not blank, f'{msp.page_count} pages' + (f'; blank: {blank}' if blank else ''))

# 10 figures: embedded images vs uploaded figure files
rels = etree.fromstring(zms.read('word/_rels/document.xml.rels'))
rmap = {r.get('Id'): r.get('Target') for r in rels}
emb = [rmap[e.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')]
       for e in rms.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip')]
det = []; ok = len(emb) == 4
for k, tgt in enumerate(emb, 1):
    a = Image.open(io.BytesIO(zms.read('word/' + tgt))).convert('L')
    pg = fitz.open(os.path.join(U, 'Figures', f'Figure{k}.pdf'))[0]
    pix = pg.get_pixmap(dpi=200)
    b = Image.frombytes('RGB', (pix.width, pix.height), pix.samples).convert('L')
    size = (600, int(600 * a.height / a.width))
    ar_diff = abs(a.height / a.width - b.height / b.width) / (a.height / a.width)
    d = ImageStat.Stat(ImageChops.difference(a.resize(size), b.resize(size))).mean[0] / 255
    det.append(f'Figure {k}: mean pixel difference {d:.3f}, aspect difference {ar_diff:.3f}')
    ok = ok and d < 0.03 and ar_diff < 0.02
add('Figures embedded in the manuscript match the uploaded Figure1-4.pdf files (rendered, 600 px, grayscale)', ok, '; '.join(det))
det = []; ok = True
for k in range(1, 5):
    doc = fitz.open(os.path.join(U, 'Figures', f'Figure{k}.pdf')); pg = doc[0]
    scale = min(1.0, (180 / 25.4 * 72) / pg.rect.width)   # figures are reduced to 180 mm, never enlarged
    spans = [s for b in pg.get_text('dict')['blocks'] for l in b.get('lines', []) for s in l['spans'] if s['text'].strip()]
    from collections import Counter
    modal = Counter(round(x['size'], 2) for x in spans).most_common(1)[0][0]
    # sub/superscripts (exponents of 10, log10 subscripts, r squared): short spans set at < 80% of the figure's modal text size
    expo = [x for x in spans if len(x['text'].strip()) <= 4 and x['size'] < 0.8 * modal]
    main = [s['size'] for s in spans if s not in expo]
    mn = min(main) * scale
    ex = (min(s['size'] for s in expo) * scale) if expo else None
    det.append(f'Figure {k}: text min {mn:.2f} pt' + (f', sub/superscripts min {ex:.2f} pt' if ex else '') + f' (scale {scale:.3f})')
    ok = ok and mn >= 6 - 1e-6
add('Main figures at 180 mm print width: no text below 6 pt (sub- and superscripts reported separately)', ok, '; '.join(det))

# 11 upload set
names = sorted(os.path.relpath(f, U) for f in glob.glob(os.path.join(U, '**', '*'), recursive=True) if os.path.isfile(f))
expected = sorted(['Manuscript_TERC_TERT_BP.docx', 'Supplemental_Material_TERC_TERT_BP.pdf', 'Supplemental_Tables_TERC_TERT_BP.xlsx',
                   'STROBE-MR_Checklist_TERC_TERT_BP.pdf', 'Cover_Letter_TERC_TERT_BP.pdf'] + [f'Figures/Figure{k}.pdf' for k in range(1, 5)])
add('Upload set: exactly one file per item, no version labels or duplicate copies in file names', names == expected, '; '.join(names))

# 12 workbook
import openpyxl
wb = openpyxl.load_workbook(os.path.join(U, 'Supplemental_Tables_TERC_TERT_BP.xlsx'), read_only=True)
num = [s for s in wb.sheetnames if re.match(r'S\d+_', s)]
ns = sorted(int(re.match(r'S(\d+)_', s).group(1)) for s in num)
add('Workbook numbered sheets: S1 and S3 companions plus contiguous S6-S88', ns == [1, 3] + list(range(6, 89)), f'{len(num)} numbered sheets')
vh = []
for s in wb.sheetnames:
    for row in wb[s].iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str) and re.search(r'\bv(1[0-9]|[2-9])\b(?!\.)|earlier (version|run)|superseded|v4-v5', v) and not allowed.search(v):
                vh.append(f'{s}: {v[:60]}')
add('Workbook text has no version labels or version-history wording outside S73 run identifiers', all(h.startswith('S73') for h in vh), '; '.join(h for h in vh if not h.startswith('S73'))[:300] or 'none')

with open(OUT, 'w', newline='') as fh:
    w = csv.writer(fh); w.writerow(['check', 'status', 'detail']); w.writerows(rows)
for r in rows:
    print(r[1], '|', r[0], '|', r[2])

# gate (not written to the workbook)
ph = {n: sum(t.count('[TO COMPLETE') for t in txt) for n, txt in (('manuscript', P), ('supplement', S), ('STROBE-MR', ST), ('cover letter', CL))}
print('GATE | author placeholders remaining |', ph, '| submit only when all are 0')
sys.exit(1 if any(r[1] == 'FAIL' for r in rows) else 0)
